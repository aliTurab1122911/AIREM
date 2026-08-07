from __future__ import annotations

import json
import os
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from scripts.document_selector import analyse_document_blocks, sections_from_visual_ranges
from scripts.openai_edit_service import build_edit_items, edit_ranges_with_openai, manual_edit_drafts
from scripts.pipeline_engine import reinsert_sections


class V19OpenAIRangeEditTests(unittest.TestCase):
    def _make_doc(self, path: Path) -> None:
        doc = Document()
        doc.add_heading("Introduction", level=1)
        p = doc.add_paragraph()
        p.add_run("The original ").bold = True
        p.add_run("sentence needs a clearer editorial revision.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Summary"
        table.cell(1, 1).text = "This table sentence can be revised manually."
        doc.save(path)

    def _ranges(self, inventory):
        paragraph = next(item for item in inventory["visual_elements"] if "sentence needs" in item["text"])
        table_text = next(item for item in inventory["visual_elements"] if "table sentence" in item["text"])
        p_start = paragraph["text"].index("sentence")
        t_start = table_text["text"].index("This")
        return [
            {
                "start_id": paragraph["visual_id"],
                "start_offset": p_start,
                "end_id": paragraph["visual_id"],
                "end_offset": len(paragraph["text"]),
                "source": "manual",
            },
            {
                "start_id": table_text["visual_id"],
                "start_offset": t_start,
                "end_id": table_text["visual_id"],
                "end_offset": len(table_text["text"]),
                "source": "manual",
            },
        ]

    def test_manual_drafts_and_direct_reinsertion_preserve_surrounding_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            output = root / "output.docx"
            self._make_doc(source)
            inventory = analyse_document_blocks(source)
            sections = sections_from_visual_ranges(source, inventory, self._ranges(inventory))
            items = build_edit_items(sections, inventory)
            drafts = manual_edit_drafts(items)["edits"]
            self.assertEqual(len(drafts), 2)
            replacements = [
                "sentence has been revised for clarity.",
                "This table sentence was revised manually.",
            ]
            result = reinsert_sections(
                source,
                {"sections": sections},
                [[value] for value in replacements],
                output,
            )
            self.assertEqual(result["replacement_count"], 2)
            reopened = Document(output)
            self.assertEqual(
                reopened.paragraphs[1].text,
                "The original sentence has been revised for clarity.",
            )
            self.assertEqual(
                reopened.tables[0].cell(1, 1).text,
                "This table sentence was revised manually.",
            )
            self.assertTrue(reopened.paragraphs[1].runs[0].bold)

    def test_openai_responses_request_uses_structured_output_and_store_false(self):
        captured = {}

        class FakeUsage:
            def model_dump(self):
                return {"input_tokens": 20, "output_tokens": 10}

        class FakeResponse:
            id = "resp_test"
            output_text = json.dumps({
                "edits": [
                    {"id": "edit_001", "revised_text": "Revised text."},
                ]
            })
            usage = FakeUsage()

        class FakeResponses:
            def create(self, **kwargs):
                captured.update(kwargs)
                return FakeResponse()

        class FakeClient:
            def __init__(self, api_key):
                captured["api_key"] = api_key
                self.responses = FakeResponses()

        fake_module = types.SimpleNamespace(OpenAI=FakeClient)
        items = [{
            "id": "edit_001",
            "source_text": "Original text.",
            "context_before": "",
            "context_after": "",
            "style": "Normal",
            "location_label": "paragraph 1",
            "visual_id": "vp_0",
            "start_offset": 0,
            "end_offset": 14,
        }]
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test-key", "OPENAI_MODEL": "gpt-5-mini"}, clear=False):
            with patch.dict(sys.modules, {"openai": fake_module}):
                result = edit_ranges_with_openai(items, "Improve clarity.")
        self.assertEqual(result["edits"][0]["revised_text"], "Revised text.")
        self.assertEqual(captured["model"], "gpt-5-mini")
        self.assertFalse(captured["store"])
        self.assertEqual(captured["text"]["format"]["type"], "json_schema")
        self.assertIn("editorial_instruction", captured["input"])


    def test_more_than_120_ranges_are_batched_and_merged_in_order(self):
        captured_inputs = []

        class FakeUsage:
            def model_dump(self):
                return {"input_tokens": 5, "output_tokens": 3, "total_tokens": 8}

        class FakeResponse:
            def __init__(self, index, edits):
                self.id = f"resp_{index}"
                self.output_text = json.dumps({"edits": edits})
                self.usage = FakeUsage()

        class FakeResponses:
            def create(self, **kwargs):
                payload = json.loads(kwargs["input"])
                captured_inputs.append(payload)
                edits = [
                    {"id": item["id"], "revised_text": f"Revised {item['id']}"}
                    for item in payload["ranges"]
                ]
                return FakeResponse(len(captured_inputs), edits)

        class FakeClient:
            def __init__(self, api_key):
                self.responses = FakeResponses()

        items = [
            {
                "id": f"edit_{index:03d}",
                "source_text": f"Original range {index}.",
                "context_before": "",
                "context_after": "",
                "style": "Normal",
                "location_label": f"paragraph {index}",
            }
            for index in range(1, 246)
        ]
        fake_module = types.SimpleNamespace(OpenAI=FakeClient)
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test-key"}, clear=False):
            with patch.dict(sys.modules, {"openai": fake_module}):
                result = edit_ranges_with_openai(items, "Improve clarity.")

        self.assertEqual(len(captured_inputs), 3)
        self.assertEqual([len(row["ranges"]) for row in captured_inputs], [120, 120, 5])
        self.assertEqual(result["batch_count"], 3)
        self.assertEqual(len(result["edits"]), 245)
        self.assertEqual(result["edits"][0]["id"], "edit_001")
        self.assertEqual(result["edits"][-1]["id"], "edit_245")
        self.assertEqual(result["usage"]["total_tokens"], 24)
        self.assertEqual(result["response_ids"], ["resp_1", "resp_2", "resp_3"])

    def test_ui_and_routes_expose_review_export_and_continue_workflow(self):
        root = Path(__file__).resolve().parents[1]
        template = (root / "templates" / "range_selector.html").read_text(encoding="utf-8")
        workspace = (root / "templates" / "workspace.html").read_text(encoding="utf-8")
        app_source = (root / "app.py").read_text(encoding="utf-8")
        self.assertIn("OpenAI edit studio", template)
        self.assertIn("Generate OpenAI drafts", template)
        self.assertIn("Manual edit only", template)
        self.assertIn("Reinsert and export DOCX", template)
        self.assertIn("Continue in AIREM", template)
        self.assertIn("toggleDraftPreview", template)
        self.assertIn("processed automatically", template)
        self.assertIn('/api/range-edit/<job_id>/export', app_source)
        self.assertIn('/api/range-edit/<job_id>/continue', app_source)
        self.assertIn("prefill_edited_texts", workspace)


if __name__ == "__main__":
    unittest.main()
