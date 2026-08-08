from __future__ import annotations

import io
import json
from pathlib import Path

import fitz
from docx import Document

import app as legacy
from processor_app import app
from scripts.document_selector import sections_from_visual_ranges


CYAN = (0.32, 0.777, 0.855)
GOLDEN = json.loads((Path(__file__).parent / "fixtures" / "turnitin_range_golden.json").read_text(encoding="utf-8"))


def _build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("The first paragraph establishes the purpose of the demonstration.")
    doc.add_paragraph(GOLDEN["paragraph_fragment"])
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Criterion"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "Mapping"
    table.cell(1, 1).text = GOLDEN["table_fragment"]
    doc.add_paragraph("The final paragraph confirms that the source document remains intact.")
    doc.save(path)


def _build_turnitin_pdf(path: Path, source_name: str) -> None:
    pdf = fitz.open()
    cover = pdf.new_page()
    cover.insert_text((72, 72), "Document Details")
    cover.insert_text((72, 94), "Submission ID")
    cover.insert_text((72, 116), "trn:test:123")
    cover.insert_text((72, 138), "File Name")
    cover.insert_text((72, 160), source_name)
    cover.insert_text((72, 182), "49 Words")
    overview = pdf.new_page()
    overview.insert_text((72, 72), "88% detected as AI")
    page = pdf.new_page()
    lines = [
        "Introduction",
        "The first paragraph establishes the purpose of the demonstration.",
        GOLDEN["paragraph_fragment"],
        "Criterion Result",
        f"Mapping {GOLDEN['table_fragment']}",
        "The final paragraph confirms that the source document remains intact.",
    ]
    y = 80
    for index, line in enumerate(lines):
        width = fitz.get_text_length(line, fontname="helv", fontsize=11)
        if index in {2, 4}:
            page.draw_rect(fitz.Rect(69, y - 11, 75 + width, y + 3), color=None, fill=CYAN, overlay=False)
        page.insert_text((72, y), line, fontsize=11, fontname="helv", overlay=True)
        y += 30
    pdf.save(path)
    pdf.close()


def _upload_document(client, source: Path):
    response = client.post(
        "/internal/v1/documents",
        data={"docx_file": (io.BytesIO(source.read_bytes()), source.name), "max_words": "4500"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 201
    return response.get_json()


def _upload_report(client, job_id: str, report: Path):
    response = client.post(
        f"/internal/v1/turnitin/{job_id}",
        data={"turnitin_pdf": (io.BytesIO(report.read_bytes()), report.name)},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    return response.get_json()


def test_turnitin_exact_and_expanded_ranges_match_golden(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "WORKSPACE", tmp_path / "workspace")
    legacy.WORKSPACE.mkdir()
    source = tmp_path / "source.docx"
    report = tmp_path / "turnitin.pdf"
    _build_docx(source)
    _build_turnitin_pdf(report, source.name)
    client = app.test_client()

    uploaded = _upload_document(client, source)
    result = _upload_report(client, uploaded["job_id"], report)
    analysis = result["analysis"]

    assert analysis["verification_status"] == GOLDEN["verification_status"]
    assert analysis["reported_ai_score"] == GOLDEN["reported_ai_score"]
    assert analysis["content_similarity"] >= GOLDEN["minimum_content_similarity"]
    assert analysis["highlight_mapping_coverage"] >= GOLDEN["minimum_highlight_mapping_coverage"]
    assert analysis["mapping_enabled"] is True
    assert analysis["exact_range_count"] >= GOLDEN["minimum_exact_ranges"]
    assert analysis["expanded_paragraph_count"] == GOLDEN["expanded_paragraph_count"]
    assert analysis["expanded_table_count"] == GOLDEN["expanded_table_count"]
    assert {item["source"] for item in analysis["exact_ranges"]} == set(GOLDEN["exact_sources"])
    assert {item["source"] for item in analysis["expanded_ranges"]} == set(GOLDEN["expanded_sources"])
    assert any(item["start_id"] == GOLDEN["paragraph_visual_id"] for item in analysis["exact_ranges"])
    assert any(item["start_id"] == GOLDEN["table_visual_id"] for item in analysis["exact_ranges"])

    inventory = uploaded["inventory"]
    exact_sections = sections_from_visual_ranges(source, inventory, analysis["exact_ranges"])
    expanded_sections = sections_from_visual_ranges(source, inventory, analysis["expanded_ranges"])
    assert any(GOLDEN["paragraph_fragment"] in section["lines"][0] for section in exact_sections)
    assert len(expanded_sections) > len(exact_sections)
    assert any(section.get("section_type") == "table_cell_paragraph_span" for section in exact_sections)


def test_manual_range_session_export_and_continue_preserve_approved_golden_text(tmp_path, monkeypatch):
    monkeypatch.setattr(legacy, "WORKSPACE", tmp_path / "workspace")
    legacy.WORKSPACE.mkdir()
    source = tmp_path / "source.docx"
    report = tmp_path / "turnitin.pdf"
    _build_docx(source)
    _build_turnitin_pdf(report, source.name)
    client = app.test_client()

    uploaded = _upload_document(client, source)
    job_id = uploaded["job_id"]
    turnitin = _upload_report(client, job_id, report)["analysis"]

    draft_response = client.post(
        f"/internal/v1/range-edit/{job_id}/draft",
        json={"visual_ranges": turnitin["exact_ranges"], "manual_only": True},
    )
    assert draft_response.status_code == 200
    draft = draft_response.get_json()
    assert draft["ok"] is True
    assert draft["mode"] == "manual"
    assert len(draft["edits"]) >= 2

    reviewed = []
    for item in draft["edits"]:
        revised = item["source_text"]
        if item["visual_id"] == GOLDEN["paragraph_visual_id"]:
            revised = GOLDEN["approved_paragraph"]
        elif item["visual_id"] == GOLDEN["table_visual_id"]:
            revised = GOLDEN["approved_table"]
        reviewed.append({"id": item["id"], "revised_text": revised})

    exported = client.post(
        f"/internal/v1/range-edit/{job_id}/export",
        json={"session_id": draft["session_id"], "edits": reviewed},
    )
    assert exported.status_code == 200
    export_body = exported.get_json()
    assert export_body["ok"] is True
    assert export_body["job_id"] == job_id
    assert export_body["replacement_count"] == len(reviewed)
    output_name = export_body["download_url"].rsplit("/", 1)[-1]
    output = legacy.WORKSPACE / job_id / "4_org_fin" / output_name
    assert output.exists()
    reopened = Document(output)
    assert reopened.paragraphs[2].text == GOLDEN["approved_paragraph"]
    assert reopened.tables[0].cell(1, 1).text == GOLDEN["approved_table"]
    assert reopened.paragraphs[1].text == "The first paragraph establishes the purpose of the demonstration."

    continued = client.post(
        f"/internal/v1/range-edit/{job_id}/continue",
        json={"session_id": draft["session_id"], "edits": reviewed},
    )
    assert continued.status_code == 200
    continue_body = continued.get_json()
    assert continue_body["selection_mode"] == "approved_range_edits"
    assert continue_body["mapping"]["selection"]["mode"] == "approved_range_edits"
    combined = "\n".join(continue_body["edited_texts"].values())
    assert GOLDEN["approved_paragraph"] in combined
    assert GOLDEN["approved_table"] in combined
    assert continue_body["chunks"] == continue_body["mapping"]["chunks"]
    assert all(
        continue_body["edited_texts"][str(chunk["chunk_number"])] == chunk["text"]
        for chunk in continue_body["chunks"]
    )

    rewritten = client.post(
        f"/internal/v1/rewrite-jobs/{job_id}/rewrite",
        json={"profile": "light", "engine": "linguistic", "style_profile": "natural_student"},
    )
    assert rewritten.status_code == 200
    rewrite_body = rewritten.get_json()
    assert rewrite_body["ok"] is True
    # The approved text is now the extraction source; any safety fallback must
    # return that reviewed source rather than the pre-review Turnitin text.
    if not rewrite_body["pass_accepted"]:
        fallback = "\n".join(rewrite_body["edited_texts"].values())
        assert GOLDEN["approved_paragraph"] in fallback
        assert GOLDEN["approved_table"] in fallback
