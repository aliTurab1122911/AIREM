from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.document_selector import analyse_document_blocks, resolve_selected_ids, sections_from_selected_blocks
from scripts.formatting_service import analyse_document_formatting, apply_document_formatting
from scripts.local_rewriter import get_profile
from scripts.rewrite_optimizer import evaluate_pass
from scripts.text_rewriter_service import rewrite_pasted_text


class V15DocumentStudioTests(unittest.TestCase):
    def _make_doc(self, path: Path) -> None:
        doc = Document()
        doc.add_heading("Cover", 0)
        doc.add_paragraph("Student Name")
        doc.add_heading("Introduction", 1)
        doc.add_paragraph("It is important to note that the system is capable of providing an explanation of the result.")
        doc.add_paragraph("First list item", style="List Bullet")
        doc.add_paragraph("Second list item", style="List Bullet")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Accuracy"
        table.cell(1, 1).text = "95%"
        doc.add_paragraph("Figure 1: System architecture")
        doc.add_heading("References", 1)
        doc.add_paragraph("Reference content")
        doc.save(path)

    def test_range_selector_supports_auto_range_and_manual(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            self._make_doc(path)
            inventory = analyse_document_blocks(path)
            automatic = resolve_selected_ids(inventory, "automatic")
            self.assertTrue(automatic)
            self.assertFalse(any(block_id.startswith("p_0") for block_id in automatic))
            manual = resolve_selected_ids(inventory, "manual", manual_ids=["p_3", "t_0"])
            self.assertEqual(manual, ["p_3", "t_0"])
            ranged = resolve_selected_ids(inventory, "range", start_order=2, end_order=6)
            self.assertTrue(ranged)
            sections = sections_from_selected_blocks(path, automatic)
            self.assertGreaterEqual(len(sections), 2)

    def test_pasted_text_rewriter_preserves_lines_and_compresses(self):
        source = (
            "It is important to note that the proposed system is capable of providing an explanation of the decision.\n\n"
            "In order to run the test, users utilise the script."
        )
        result = rewrite_pasted_text(source, "rewrite_compress")
        self.assertTrue(result["ok"])
        self.assertIn("\n\n", result["rewritten_text"])
        self.assertLess(result["rewritten_words"], result["original_words"])
        self.assertIn("can explain", result["rewritten_text"])

    def test_formatting_service_adds_fields_and_reopens(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "source.docx"
            out = Path(tmp) / "formatted.docx"
            self._make_doc(src)
            report = analyse_document_formatting(src)
            self.assertEqual(report["manual_caption_count"], 1)
            result = apply_document_formatting(src, out, {
                "body_font": "Arial",
                "body_size": "11",
                "body_color": "#000000",
                "heading_font": "Arial",
                "heading_color": "#1F4E79",
                "convert_captions": True,
                "page_numbers": True,
                "add_toc": True,
                "add_lof": True,
                "add_lot": True,
                "repeat_header": True,
                "bold_table_header": True,
            })
            self.assertEqual(result["captions"]["Figure"], 1)
            reopened = Document(out)
            self.assertGreater(len(reopened.paragraphs), 1)

    def test_balanced_and_permissive_reduce_rollback_strictness(self):
        original = {"word_count": 1000, "ai_style_score": 20.0}
        source = {"word_count": 1000, "ai_style_score": 15.0}
        slight_tradeoff = {"word_count": 1000, "ai_style_score": 15.1}
        strict = evaluate_pass(original, source, slight_tradeoff, get_profile("natural"), "strict")
        balanced = evaluate_pass(original, source, slight_tradeoff, get_profile("natural"), "balanced")
        permissive = evaluate_pass(original, source, slight_tradeoff, get_profile("natural"), "permissive")
        self.assertFalse(strict["accepted"])
        self.assertTrue(balanced["accepted"] or permissive["accepted"])
        self.assertTrue(permissive["accepted"])


if __name__ == "__main__":
    unittest.main()
