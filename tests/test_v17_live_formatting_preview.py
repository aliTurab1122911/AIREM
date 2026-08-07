from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.document_selector import analyse_document_blocks
from scripts.formatting_service import apply_document_formatting


class V17LiveFormattingPreviewTests(unittest.TestCase):
    def _make_doc(self, path: Path) -> None:
        doc = Document()
        title = doc.add_paragraph(style="Title")
        title_run = title.add_run("Preview Test")
        title_run.font.name = "Times New Roman"
        body = doc.add_paragraph()
        body_run = body.add_run("Body text with direct formatting.")
        body_run.font.name = "Times New Roman"
        doc.add_heading("Section Heading", level=1)
        doc.add_paragraph("Figure 1: Preview diagram")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header A"
        table.cell(0, 1).text = "Header B"
        table.cell(1, 0).text = "Value A"
        table.cell(1, 1).text = "Value B"
        doc.save(path)

    def test_formatting_workspace_contains_live_word_preview_controls(self):
        root = Path(__file__).resolve().parents[1]
        template = (root / "templates" / "formatting_workspace.html").read_text(encoding="utf-8")
        app_source = (root / "app.py").read_text(encoding="utf-8")
        self.assertIn('id="wordPage"', template)
        self.assertIn('id="originalView"', template)
        self.assertIn('id="liveView"', template)
        self.assertIn('function updatePreview()', template)
        self.assertIn('id="tocPreview"', template)
        self.assertIn('preview = analyse_document_blocks(original_path)', app_source)
        self.assertIn('"preview_path": str(preview_path)', app_source)

    def test_preview_inventory_contains_renderable_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            self._make_doc(source)
            inventory = analyse_document_blocks(source)
            self.assertGreaterEqual(inventory["paragraph_count"], 4)
            self.assertEqual(inventory["table_count"], 1)
            self.assertTrue(any(block["type"] == "paragraph" and block["visual"]["runs"] for block in inventory["blocks"]))
            self.assertTrue(any(block["type"] == "table" and block["rendered_rows"] for block in inventory["blocks"]))

    def test_generated_docx_matches_preview_typography_override(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            output = Path(tmp) / "formatted.docx"
            self._make_doc(source)
            result = apply_document_formatting(source, output, {
                "body_font": "Arial",
                "body_size": "12",
                "body_color": "#123456",
                "heading_font": "Calibri",
                "heading_color": "#654321",
                "title_size": "24",
                "h1_size": "18",
                "h2_size": "15",
                "h3_size": "13",
                "caption_font": "Verdana",
                "caption_size": "9",
                "caption_color": "#222222",
                "table_font": "Tahoma",
                "table_size": "10",
                "table_border_color": "#808080",
                "table_border_size": "4",
            })
            self.assertTrue(result["direct_run_typography_applied"])
            reopened = Document(output)
            self.assertEqual(reopened.paragraphs[1].runs[0].font.name, "Arial")
            self.assertEqual(round(reopened.paragraphs[1].runs[0].font.size.pt, 1), 12.0)
            self.assertEqual(reopened.paragraphs[2].runs[0].font.name, "Calibri")
            self.assertEqual(reopened.tables[0].cell(1, 0).paragraphs[0].runs[0].font.name, "Tahoma")


if __name__ == "__main__":
    unittest.main()
