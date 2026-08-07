import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document

from scripts.document_selector import analyse_document_blocks, sections_from_visual_ranges
from scripts.turnitin_mapper import analyse_turnitin_report


CYAN = (0.32, 0.777, 0.855)


class V18TurnitinMappingTests(unittest.TestCase):
    def _build_docx(self, path: Path) -> None:
        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("The first paragraph establishes the purpose of the demonstration.")
        doc.add_paragraph("This exact paragraph should be highlighted and mapped into the Word preview.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Criterion"
        table.cell(0, 1).text = "Result"
        table.cell(1, 0).text = "Mapping"
        table.cell(1, 1).text = "This highlighted table cell should expand to the complete table."
        doc.add_paragraph("The final paragraph confirms that the source document remains intact.")
        doc.save(path)

    def _build_report(self, path: Path, source_name: str, mismatch: bool = False) -> None:
        pdf = fitz.open()
        cover = pdf.new_page()
        cover.insert_text((72, 72), "Document Details")
        cover.insert_text((72, 94), "Submission ID")
        cover.insert_text((72, 116), "trn:test:123")
        cover.insert_text((72, 138), "File Name")
        cover.insert_text((72, 160), source_name)
        cover.insert_text((72, 182), "49 Words")
        overview = pdf.new_page()
        overview.insert_text((72, 72), "88% detected as AI")
        page = pdf.new_page()
        lines = [
            "Introduction",
            "Completely unrelated report content that cannot match the source." if mismatch else "The first paragraph establishes the purpose of the demonstration.",
            "This exact paragraph should be highlighted and mapped into the Word preview.",
            "Criterion Result",
            "Mapping This highlighted table cell should expand to the complete table.",
            "The final paragraph confirms that the source document remains intact.",
        ]
        y = 80
        for index, line in enumerate(lines):
            width = fitz.get_text_length(line, fontname="helv", fontsize=11)
            if not mismatch and index in {2, 4}:
                page.draw_rect(fitz.Rect(69, y - 11, 75 + width, y + 3), color=None, fill=CYAN, overlay=False)
            page.insert_text((72, y), line, fontsize=11, fontname="helv", overlay=True)
            y += 30
        pdf.save(path)
        pdf.close()

    def test_verified_report_maps_exact_and_expanded_ranges(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            docx = root / "source.docx"
            report = root / "turnitin.pdf"
            self._build_docx(docx)
            self._build_report(report, "source.docx")
            inventory = analyse_document_blocks(docx)
            result = analyse_turnitin_report(docx, "source.docx", inventory, report)

            self.assertEqual(result["verification_status"], "verified")
            self.assertTrue(result["mapping_enabled"])
            self.assertGreaterEqual(result["highlight_mapping_coverage"], 0.95)
            self.assertGreaterEqual(result["exact_range_count"], 2)
            self.assertEqual(result["expanded_table_count"], 1)
            self.assertGreaterEqual(result["expanded_paragraph_count"], 1)
            self.assertEqual(result["reported_ai_score"], 88)

            exact_sections = sections_from_visual_ranges(docx, inventory, result["exact_ranges"])
            expanded_sections = sections_from_visual_ranges(docx, inventory, result["expanded_ranges"])
            self.assertTrue(any("exact paragraph" in section["lines"][0] for section in exact_sections))
            self.assertGreater(len(expanded_sections), len(exact_sections))

    def test_mismatched_report_disables_automatic_mapping(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            docx = root / "source.docx"
            report = root / "turnitin.pdf"
            self._build_docx(docx)
            self._build_report(report, "different.docx", mismatch=True)
            inventory = analyse_document_blocks(docx)
            result = analyse_turnitin_report(docx, "source.docx", inventory, report)
            self.assertIn(result["verification_status"], {"review", "mismatch"})
            self.assertFalse(result["mapping_enabled"])
            self.assertFalse(result["filename_match"])

    def test_range_selector_exposes_turnitin_workflow(self):
        template = (Path(__file__).resolve().parents[1] / "templates" / "range_selector.html").read_text(encoding="utf-8")
        self.assertIn("Map Turnitin highlights to DOCX", template)
        self.assertIn("Expand all to full paragraphs/tables", template)
        self.assertIn("Turnitin report preview", template)
        self.assertIn("applyTurnitinRanges", template)


if __name__ == "__main__":
    unittest.main()
