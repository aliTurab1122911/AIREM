from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.detection_service import detect_docx, detect_text
from scripts.document_selector import analyse_document_blocks, sections_from_visual_ranges
from scripts.docx_pipeline_common import sections_to_text
from scripts.pipeline_engine import create_extraction_package_from_sections, reinsert_sections, validate_edited_chunks


class V16DetectionAndVisualRangeTests(unittest.TestCase):
    def test_detection_v3_is_bounded_and_exposes_components(self):
        text = (
            "Furthermore, it is important to note that the proposed framework provides a comprehensive "
            "and systematic mechanism for evaluating the final classification decision. "
            "Moreover, the proposed framework provides a comprehensive and systematic mechanism for "
            "evaluating the final classification decision. Therefore, the findings indicate that the "
            "framework provides a comprehensive and systematic mechanism for evaluation."
        )
        report = detect_text(text)
        self.assertGreaterEqual(report["score"], 0)
        self.assertLessEqual(report["score"], 100)
        self.assertEqual(report["method"], "transparent_style_pattern_ensemble_v3")
        self.assertIn("structural_regularity", report["dimensions"])
        self.assertIn("highest_risk_passages", report)
        self.assertIn("does not prove", report["disclaimer"])

    def test_docx_detection_reads_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            doc = Document()
            doc.add_paragraph("This paragraph contains enough text to support a stable passage-level diagnostic result for testing.")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Heading"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "A repeated and formal table description; therefore, the result is recorded."
            table.cell(1, 1).text = "Recorded value"
            doc.save(path)
            report = detect_docx(path)
            self.assertEqual(report["document"]["table_count"], 1)
            self.assertGreater(report["word_count"], 0)
            self.assertTrue(report["highest_risk_passages"])

    def test_multiple_visual_ranges_reinsert_only_selected_spans(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            output = root / "output.docx"
            mapping_path = root / "map.json"
            extracts = root / "extracts"
            log = root / "log.json"

            doc = Document()
            p1 = doc.add_paragraph()
            p1.add_run("Alpha ")
            italic = p1.add_run("beta")
            italic.italic = True
            p1.add_run(" gamma delta.")
            doc.add_paragraph("Second paragraph remains mostly unchanged.")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Cell text remains visible."
            doc.save(source)

            inventory = analyse_document_blocks(source)
            ranges = [
                {"start_id": "vp_0", "start_offset": 6, "end_id": "vp_0", "end_offset": 10},
                {"start_id": "vp_1", "start_offset": 0, "end_id": "vp_1", "end_offset": 6},
                {"start_id": "vc_0_0_0_0", "start_offset": 0, "end_id": "vc_0_0_0_0", "end_offset": 4},
            ]
            sections = sections_from_visual_ranges(source, inventory, ranges)
            self.assertEqual(len(sections), 3)
            mapping = create_extraction_package_from_sections(
                source, sections, extracts, mapping_path, max_words=4500,
                selection_metadata={"mode": "visual", "visual_ranges": ranges},
            )
            edited = {1: sections_to_text([["BETA"], ["SECOND"], ["CELL"]])}
            validation = validate_edited_chunks(mapping, edited)
            self.assertTrue(validation["valid"], validation)
            result = reinsert_sections(source, mapping, validation["parsed_sections"], output, log)
            self.assertEqual(result["replacement_count"], 3)

            reopened = Document(output)
            self.assertEqual(reopened.paragraphs[0].text, "Alpha BETA gamma delta.")
            self.assertEqual(reopened.paragraphs[1].text, "SECOND paragraph remains mostly unchanged.")
            self.assertEqual(reopened.tables[0].cell(0, 0).text, "CELL text remains visible.")
            beta_runs = [run for run in reopened.paragraphs[0].runs if "BETA" in run.text]
            self.assertTrue(beta_runs)
            self.assertTrue(beta_runs[0].italic)

    def test_rewrite_cycle_limit_removed_from_server_and_workspace(self):
        root = Path(__file__).resolve().parents[1]
        app_source = (root / "app.py").read_text(encoding="utf-8")
        workspace = (root / "templates" / "workspace.html").read_text(encoding="utf-8")
        self.assertNotIn("maximum_cycle_limit_reached", app_source)
        self.assertNotIn('id="max_cycles"', workspace)
        self.assertIn("unlimited_rewrite_cycles", app_source)


if __name__ == "__main__":
    unittest.main()
