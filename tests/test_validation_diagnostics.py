import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.docx_pipeline_common import parse_extract_text, sections_to_text
from scripts.pipeline_engine import (
    create_extraction_package,
    extract_sections_from_docx,
    validate_edited_chunks,
)


class ValidationDiagnosticsTests(unittest.TestCase):
    @staticmethod
    def mapping(section_count=2):
        sections = []
        for index in range(section_count):
            sections.append({
                "section_type": "paragraph_group",
                "items": [{"kind": "paragraph", "original_text": f"Line {index + 1}."}],
            })
        return {
            "section_count": section_count,
            "sections": sections,
            "chunks": [{
                "chunk_number": 1,
                "section_indices": list(range(section_count)),
            }],
        }

    def test_empty_chunk_has_specific_diagnostic(self):
        report = validate_edited_chunks(self.mapping(), {1: ""})
        self.assertFalse(report["valid"])
        self.assertEqual(report["section_count_actual"], 0)
        self.assertEqual(report["issues"][0]["code"], "empty_edited_chunk")
        self.assertIn("not a list-formatting error", report["issues"][0]["message"])
        chunk = report["chunks"][0]
        self.assertTrue(chunk["input_present"])
        self.assertEqual(chunk["input_character_count"], 0)
        self.assertEqual(chunk["divider_count"], 0)

    def test_missing_chunk_key_is_reported_as_empty(self):
        report = validate_edited_chunks(self.mapping(), {})
        self.assertEqual(report["issues"][0]["code"], "empty_edited_chunk")
        self.assertFalse(report["chunks"][0]["input_present"])

    def test_divider_only_chunk_has_specific_diagnostic(self):
        report = validate_edited_chunks(self.mapping(), {1: "|sec|\n|sec|"})
        self.assertEqual(report["issues"][0]["code"], "divider_only_chunk")
        self.assertEqual(report["chunks"][0]["divider_count"], 2)

    def test_nonempty_text_without_markers_reports_missing_dividers(self):
        report = validate_edited_chunks(self.mapping(), {1: "First section.\nSecond section."})
        self.assertEqual(report["section_count_actual"], 1)
        self.assertEqual(report["issues"][0]["code"], "missing_section_dividers")

    def test_generated_markers_validate_word_lists(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "lists.docx"
            extract_dir = tmp_path / "extract"
            map_path = tmp_path / "map.json"

            doc = Document()
            doc.add_heading("Executive Summary", level=1)
            doc.add_paragraph("Introductory prose.")
            doc.add_paragraph("First bullet item.", style="List Bullet")
            doc.add_paragraph("Second bullet item.", style="List Bullet")
            doc.add_paragraph("First numbered item.", style="List Number")
            doc.add_paragraph("Second numbered item.", style="List Number")
            doc.add_heading("References", level=1)
            doc.save(source)

            sections = extract_sections_from_docx(source)
            self.assertEqual(len(sections), 1)
            self.assertEqual(len(sections[0]["lines"]), 5)

            mapping = create_extraction_package(source, extract_dir, map_path)
            generated = mapping["chunks"][0]["text"]
            self.assertEqual(parse_extract_text(generated), [sections[0]["lines"]])

            report = validate_edited_chunks(mapping, {1: generated})
            self.assertTrue(report["valid"], report["issues"])
            self.assertEqual(report["section_count_actual"], 1)
            self.assertGreaterEqual(report["chunks"][0]["divider_count"], 2)


if __name__ == "__main__":
    unittest.main()
