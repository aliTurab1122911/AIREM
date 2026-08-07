from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

SEC = "|sec|"
WORD_RE = re.compile(r"\b[\w'-]+\b", flags=re.UNICODE)


def iter_block_items(parent):
    """Yield Paragraph and Table objects in document order for a document or cell."""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise TypeError(f"Unsupported parent type: {type(parent)}")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def norm_text(text: str) -> str:
    text = text.replace("\r", "")
    lines = [ln.rstrip() for ln in text.split("\n")]
    return "\n".join(lines).strip()


def word_count(text: str) -> int:
    return len(WORD_RE.findall(text or ""))


def section_word_count(lines: List[str]) -> int:
    return sum(word_count(line) for line in lines)


def one_line_text(text: str) -> str:
    """Normalize any Word paragraph/cell text into one reinsertion-safe line."""
    return " ".join(norm_text(text).splitlines()).strip()


def para_text(paragraph: Paragraph) -> str:
    return one_line_text(paragraph.text)


def cell_text(cell: _Cell) -> str:
    parts = [one_line_text(p.text) for p in cell.paragraphs]
    parts = [p for p in parts if p]
    return " ".join(parts).strip()


def is_heading(paragraph: Paragraph) -> bool:
    name = (paragraph.style.name or "").lower()
    return name.startswith("heading")


def is_caption_or_label(text: str) -> bool:
    t = text.strip().lower()
    return bool(re.match(r"^(figure|table)\s+\d+\s*:", t))


def replace_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    """Replace visible text while preserving paragraph style and first-run formatting as much as possible."""
    new_text = norm_text(new_text)
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)




def replace_paragraph_span(paragraph: Paragraph, start_offset: int, end_offset: int, new_text: str) -> None:
    """Replace a character span while preserving surrounding run formatting.

    The replacement inherits the formatting of the run containing the start of
    the selected span. Offsets are measured against ``paragraph.text``.
    """
    start = max(0, int(start_offset))
    end = max(start, int(end_offset))
    source = paragraph.text or ""
    start = min(start, len(source))
    end = min(end, len(source))
    replacement = norm_text(new_text)

    if not paragraph.runs:
        paragraph.add_run(source[:start] + replacement + source[end:])
        return

    boundaries = []
    cursor = 0
    for run in paragraph.runs:
        run_start = cursor
        cursor += len(run.text or "")
        boundaries.append((run, run_start, cursor))

    start_index = None
    end_index = None
    for index, (_, run_start, run_end) in enumerate(boundaries):
        if start_index is None and (start < run_end or (start == len(source) and index == len(boundaries) - 1)):
            start_index = index
        if end <= run_end:
            end_index = index
            break
    if start_index is None:
        start_index = len(boundaries) - 1
    if end_index is None:
        end_index = len(boundaries) - 1

    start_run, start_run_start, _ = boundaries[start_index]
    end_run, end_run_start, _ = boundaries[end_index]
    prefix = (start_run.text or "")[: max(0, start - start_run_start)]
    suffix = (end_run.text or "")[max(0, end - end_run_start):]

    if start_index == end_index:
        start_run.text = prefix + replacement + suffix
        return

    start_run.text = prefix + replacement
    for index in range(start_index + 1, end_index):
        boundaries[index][0].text = ""
    end_run.text = suffix

def replace_cell_text(cell: _Cell, new_text: str) -> None:
    new_text = norm_text(new_text)
    if not cell.paragraphs:
        cell.add_paragraph(new_text)
        return
    replace_paragraph_text(cell.paragraphs[0], new_text)
    for p in cell.paragraphs[1:]:
        replace_paragraph_text(p, "")


def parse_extract_text(raw_text: str) -> List[List[str]]:
    """Parse pasted extract text. Every block between |sec| markers is a section.

    Empty lines are ignored so copy/paste from Word, browser textareas, and external editors remains stable.
    """
    sections: List[List[str]] = []
    current: List[str] = []
    for raw_line in (raw_text or "").replace("\r", "").split("\n"):
        text = norm_text(raw_line)
        if not text:
            continue
        if text == SEC:
            if current:
                sections.append(current)
                current = []
            continue
        current.append(text)
    if current:
        sections.append(current)
    return sections


def read_extract_sections(docx_path: str | Path) -> List[List[str]]:
    doc = Document(str(docx_path))
    text = "\n".join(norm_text(p.text) for p in doc.paragraphs if norm_text(p.text))
    return parse_extract_text(text)


def sections_to_text(sections: List[List[str]], include_final_sec: bool = True) -> str:
    """Return copy-safe text that starts with |sec| and optionally ends with |sec|."""
    parts: List[str] = [SEC]
    for lines in sections:
        parts.extend(lines)
        parts.append(SEC)
    if not include_final_sec and parts and parts[-1] == SEC:
        parts.pop()
    return "\n".join(parts)


def write_extract_docx(sections: List[List[str]], output_path: str | Path) -> None:
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for line in sections_to_text(sections).split("\n"):
        doc.add_paragraph(line)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def load_map(path: str | Path) -> Dict[str, Any]:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_json(data: Any, path: str | Path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
