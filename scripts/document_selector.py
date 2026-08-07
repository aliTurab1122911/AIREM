from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Mapping, Sequence

from docx import Document

from .docx_pipeline_common import (
    cell_text,
    is_caption_or_label,
    is_heading,
    iter_block_items,
    para_text,
    section_word_count,
    word_count,
)
from .pipeline_engine import (
    _is_end_heading,
    _is_front_matter_heading,
    _is_start_heading,
    should_skip_table,
)


def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip().lower()
    if not name.startswith("heading"):
        return None
    digits = "".join(ch for ch in name if ch.isdigit())
    return int(digits) if digits else 1


def _alignment_name(value: Any) -> str:
    raw = str(value).lower()
    if "center" in raw or raw == "1":
        return "center"
    if "right" in raw or raw == "2":
        return "right"
    if "justify" in raw or raw in {"3", "4"}:
        return "justify"
    return "left"


def _colour_hex(run: Any) -> str | None:
    try:
        rgb = run.font.color.rgb
        return str(rgb) if rgb else None
    except Exception:
        return None


def _effective_font(run: Any, paragraph: Any) -> tuple[str | None, float | None]:
    name = run.font.name
    size = run.font.size.pt if run.font.size else None
    try:
        if not name and paragraph.style and paragraph.style.font:
            name = paragraph.style.font.name
        if size is None and paragraph.style and paragraph.style.font and paragraph.style.font.size:
            size = paragraph.style.font.size.pt
    except Exception:
        pass
    return name, round(float(size), 2) if size is not None else None


def _run_payload(run: Any, paragraph: Any) -> Dict[str, Any]:
    name, size = _effective_font(run, paragraph)
    return {
        "text": run.text,
        "bold": bool(run.bold),
        "italic": bool(run.italic),
        "underline": bool(run.underline),
        "font_name": name,
        "font_size_pt": size,
        "colour": _colour_hex(run),
    }


def _has_numbering(paragraph: Any) -> bool:
    try:
        ppr = paragraph._p.pPr
        return bool(ppr is not None and ppr.numPr is not None)
    except Exception:
        return False


def _paragraph_visual_payload(paragraph: Any, visual_id: str, location: Dict[str, Any], order: int) -> Dict[str, Any]:
    text = paragraph.text or ""
    runs = [_run_payload(run, paragraph) for run in paragraph.runs]
    if not runs and text:
        runs = [{
            "text": text,
            "bold": False,
            "italic": False,
            "underline": False,
            "font_name": None,
            "font_size_pt": None,
            "colour": None,
        }]
    style_name = paragraph.style.name if paragraph.style else "Normal"
    left_indent = paragraph.paragraph_format.left_indent
    first_indent = paragraph.paragraph_format.first_line_indent
    return {
        "visual_id": visual_id,
        "order": order,
        "text": text,
        "word_count": word_count(text),
        "runs": runs,
        "style": style_name,
        "alignment": _alignment_name(paragraph.alignment),
        "left_indent_pt": round(left_indent.pt, 2) if left_indent else 0,
        "first_line_indent_pt": round(first_indent.pt, 2) if first_indent else 0,
        "list_item": _has_numbering(paragraph) or "list" in style_name.lower(),
        "location": location,
    }


def _page_settings(doc: Any) -> Dict[str, Any]:
    if not doc.sections:
        return {}
    section = doc.sections[0]
    return {
        "page_width_pt": round(section.page_width.pt, 2) if section.page_width else 612,
        "page_height_pt": round(section.page_height.pt, 2) if section.page_height else 792,
        "margin_top_pt": round(section.top_margin.pt, 2) if section.top_margin else 72,
        "margin_bottom_pt": round(section.bottom_margin.pt, 2) if section.bottom_margin else 72,
        "margin_left_pt": round(section.left_margin.pt, 2) if section.left_margin else 72,
        "margin_right_pt": round(section.right_margin.pt, 2) if section.right_margin else 72,
    }


def analyse_document_blocks(input_path: str | Path) -> Dict[str, Any]:
    """Return an ordered inventory plus a browser-renderable Word-style preview."""
    doc = Document(str(input_path))
    blocks: List[Dict[str, Any]] = []
    visual_elements: List[Dict[str, Any]] = []
    heading_stack: Dict[int, str] = {}
    paragraph_index = -1
    table_index = -1
    visual_order = 0
    in_default_body = False
    ended_default_body = False

    for order, block in enumerate(iter_block_items(doc)):
        if block.__class__.__name__ == "Paragraph":
            paragraph_index += 1
            text = para_text(block)
            style = block.style.name or "Normal"
            heading = is_heading(block)
            level = _heading_level(style)
            caption = is_caption_or_label(text)

            if heading and text:
                lvl = level or 1
                heading_stack[lvl] = text
                for key in list(heading_stack):
                    if key > lvl:
                        del heading_stack[key]

            if _is_start_heading(text):
                in_default_body = True
            if _is_end_heading(text) and in_default_body:
                ended_default_body = True

            default_selected = bool(
                in_default_body
                and not ended_default_body
                and text
                and not heading
                and not caption
                and not _is_front_matter_heading(text)
            )
            visual = _paragraph_visual_payload(
                block,
                visual_id=f"vp_{paragraph_index}",
                location={"kind": "paragraph", "body_paragraph_index": paragraph_index},
                order=visual_order,
            )
            visual_order += 1
            visual_elements.append(visual)
            blocks.append({
                "block_id": f"p_{paragraph_index}",
                "order": order,
                "type": "paragraph",
                "paragraph_index": paragraph_index,
                "text": text,
                "preview": text[:180] or "[Blank paragraph]",
                "style": style,
                "heading_level": level,
                "heading_path": [heading_stack[k] for k in sorted(heading_stack)],
                "is_heading": heading,
                "is_caption": caption,
                "word_count": word_count(text),
                "default_selected": default_selected,
                "visual": visual,
            })
        else:
            table_index += 1
            table = block
            rows = len(table.rows)
            cols = len(table.columns) if table.rows else 0
            headers = [cell_text(cell) for cell in table.rows[0].cells] if table.rows else []
            preview_cells: List[str] = []
            rendered_rows: List[Dict[str, Any]] = []
            for row_index, row in enumerate(table.rows):
                rendered_cells: List[Dict[str, Any]] = []
                if row_index < 2:
                    preview_cells.extend(cell_text(cell) for cell in row.cells)
                for column_index, cell in enumerate(row.cells):
                    paragraph_payloads: List[Dict[str, Any]] = []
                    for cell_paragraph_index, paragraph in enumerate(cell.paragraphs):
                        visual = _paragraph_visual_payload(
                            paragraph,
                            visual_id=f"vc_{table_index}_{row_index}_{column_index}_{cell_paragraph_index}",
                            location={
                                "kind": "table_cell_paragraph",
                                "table_index": table_index,
                                "row_index": row_index,
                                "col_index": column_index,
                                "cell_paragraph_index": cell_paragraph_index,
                            },
                            order=visual_order,
                        )
                        visual_order += 1
                        visual_elements.append(visual)
                        paragraph_payloads.append(visual)
                    rendered_cells.append({
                        "row_index": row_index,
                        "column_index": column_index,
                        "paragraphs": paragraph_payloads,
                    })
                rendered_rows.append({"row_index": row_index, "cells": rendered_cells})
            preview = " | ".join(x for x in preview_cells if x)[:180] or "[Empty table]"
            table_words = sum(word_count(cell_text(cell)) for row in table.rows for cell in row.cells)
            default_selected = bool(in_default_body and not ended_default_body and rows > 1 and not should_skip_table(headers))
            blocks.append({
                "block_id": f"t_{table_index}",
                "order": order,
                "type": "table",
                "table_index": table_index,
                "text": preview,
                "preview": preview,
                "style": getattr(table.style, "name", "Table") if table.style else "Table",
                "heading_level": None,
                "heading_path": [heading_stack[k] for k in sorted(heading_stack)],
                "is_heading": False,
                "is_caption": False,
                "rows": rows,
                "columns": cols,
                "headers": headers,
                "word_count": table_words,
                "default_selected": default_selected,
                "rendered_rows": rendered_rows,
            })

    return {
        "blocks": blocks,
        "visual_elements": visual_elements,
        "visual_element_count": len(visual_elements),
        "block_count": len(blocks),
        "paragraph_count": sum(1 for block in blocks if block["type"] == "paragraph"),
        "table_count": sum(1 for block in blocks if block["type"] == "table"),
        "default_selected_count": sum(1 for block in blocks if block["default_selected"]),
        "word_count": sum(int(block["word_count"]) for block in blocks),
        "page": _page_settings(doc),
    }


def resolve_selected_ids(
    inventory: Dict[str, Any],
    mode: str,
    manual_ids: Sequence[str] | None = None,
    start_order: int | None = None,
    end_order: int | None = None,
    include_headings: bool = False,
    include_captions: bool = False,
) -> List[str]:
    blocks = list(inventory.get("blocks", []))
    mode = (mode or "automatic").lower()
    if mode == "manual":
        requested = set(manual_ids or [])
        selected = [block for block in blocks if block["block_id"] in requested]
    elif mode == "range":
        start = 0 if start_order is None else int(start_order)
        end = len(blocks) - 1 if end_order is None else int(end_order)
        if start > end:
            start, end = end, start
        selected = [block for block in blocks if start <= int(block["order"]) <= end]
    else:
        selected = [block for block in blocks if block.get("default_selected")]

    result: List[str] = []
    for block in selected:
        if block["type"] == "paragraph":
            if block.get("is_heading") and not include_headings:
                continue
            if block.get("is_caption") and not include_captions:
                continue
            if not str(block.get("text") or "").strip():
                continue
        result.append(block["block_id"])
    return result


def sections_from_selected_blocks(
    input_path: str | Path,
    selected_ids: Sequence[str],
    include_table_headers: bool = False,
) -> List[Dict[str, Any]]:
    selected = set(selected_ids)
    doc = Document(str(input_path))
    sections: List[Dict[str, Any]] = []
    paragraph_index = -1
    table_index = -1
    paragraph_items: List[Dict[str, Any]] = []
    paragraph_lines: List[str] = []
    paragraph_block_ids: List[str] = []

    def flush_paragraphs() -> None:
        nonlocal paragraph_items, paragraph_lines, paragraph_block_ids
        if paragraph_lines:
            sections.append({
                "section_type": "paragraph_group",
                "items": paragraph_items,
                "lines": paragraph_lines,
                "word_count": section_word_count(paragraph_lines),
                "source_block_ids": paragraph_block_ids,
            })
        paragraph_items = []
        paragraph_lines = []
        paragraph_block_ids = []

    for block in iter_block_items(doc):
        if block.__class__.__name__ == "Paragraph":
            paragraph_index += 1
            block_id = f"p_{paragraph_index}"
            text = para_text(block)
            if block_id not in selected or not text:
                flush_paragraphs()
                continue
            paragraph_lines.append(text)
            paragraph_block_ids.append(block_id)
            paragraph_items.append({
                "kind": "paragraph",
                "body_paragraph_index": paragraph_index,
                "original_text": text,
                "style": block.style.name,
            })
        else:
            flush_paragraphs()
            table_index += 1
            block_id = f"t_{table_index}"
            if block_id not in selected:
                continue
            table = block
            if not table.rows:
                continue
            headers = [cell_text(cell) for cell in table.rows[0].cells]
            first_row = 0 if include_table_headers else 1
            for row_idx in range(first_row, len(table.rows)):
                lines: List[str] = []
                items: List[Dict[str, Any]] = []
                for col_idx in range(len(table.columns)):
                    text = cell_text(table.cell(row_idx, col_idx)) or ""
                    lines.append(text)
                    items.append({
                        "kind": "table_cell",
                        "table_index": table_index,
                        "row_index": row_idx,
                        "col_index": col_idx,
                        "original_text": text,
                    })
                if any(line.strip() for line in lines):
                    sections.append({
                        "section_type": "table_row",
                        "table_index": table_index,
                        "row_index": row_idx,
                        "headers": headers,
                        "items": items,
                        "lines": lines,
                        "word_count": section_word_count(lines),
                        "source_block_id": block_id,
                    })
    flush_paragraphs()
    return sections


def _normalise_visual_ranges(inventory: Mapping[str, Any], ranges: Sequence[Mapping[str, Any]]) -> Dict[str, List[tuple[int, int]]]:
    elements = list(inventory.get("visual_elements", []))
    by_id = {str(item["visual_id"]): item for item in elements}
    order = {str(item["visual_id"]): int(item["order"]) for item in elements}
    intervals: Dict[str, List[tuple[int, int]]] = {}

    for raw in ranges:
        start_id = str(raw.get("start_id") or "")
        end_id = str(raw.get("end_id") or "")
        if start_id not in by_id or end_id not in by_id:
            continue
        try:
            start_offset = int(raw.get("start_offset") or 0)
            end_offset = int(raw.get("end_offset") or 0)
        except (TypeError, ValueError):
            continue
        start_key = (order[start_id], start_offset)
        end_key = (order[end_id], end_offset)
        if start_key > end_key:
            start_id, end_id = end_id, start_id
            start_offset, end_offset = end_offset, start_offset
        start_order = order[start_id]
        end_order = order[end_id]
        for element in elements:
            element_id = str(element["visual_id"])
            element_order = int(element["order"])
            if not (start_order <= element_order <= end_order):
                continue
            text = str(element.get("text") or "")
            lower = start_offset if element_id == start_id else 0
            upper = end_offset if element_id == end_id else len(text)
            lower = max(0, min(lower, len(text)))
            upper = max(0, min(upper, len(text)))
            if lower > upper:
                lower, upper = upper, lower
            if upper <= lower:
                continue
            fragment = text[lower:upper]
            if not fragment.strip():
                continue
            # Keep surrounding spaces outside the editable span so a rewritten
            # fragment cannot accidentally join neighbouring words.
            leading = len(fragment) - len(fragment.lstrip())
            trailing = len(fragment) - len(fragment.rstrip())
            lower += leading
            upper -= trailing
            if upper <= lower:
                continue
            intervals.setdefault(element_id, []).append((lower, upper))

    merged: Dict[str, List[tuple[int, int]]] = {}
    for element_id, values in intervals.items():
        ordered = sorted(values)
        result: List[list[int]] = []
        for start, end in ordered:
            if result and start <= result[-1][1]:
                result[-1][1] = max(result[-1][1], end)
            else:
                result.append([start, end])
        merged[element_id] = [(start, end) for start, end in result]
    return merged


def sections_from_visual_ranges(
    input_path: str | Path,
    inventory: Mapping[str, Any],
    ranges: Sequence[Mapping[str, Any]],
) -> List[Dict[str, Any]]:
    """Convert one or more browser text selections into reinsertable text spans."""
    del input_path  # Mapping is already based on the uploaded immutable source.
    intervals = _normalise_visual_ranges(inventory, ranges)
    sections: List[Dict[str, Any]] = []
    for element in sorted(inventory.get("visual_elements", []), key=lambda item: int(item["order"])):
        element_id = str(element["visual_id"])
        if element_id not in intervals:
            continue
        text = str(element.get("text") or "")
        location = dict(element.get("location") or {})
        for span_index, (start, end) in enumerate(intervals[element_id], start=1):
            fragment = text[start:end]
            if not fragment.strip():
                continue
            if location.get("kind") == "paragraph":
                item = {
                    "kind": "paragraph_span",
                    "body_paragraph_index": int(location["body_paragraph_index"]),
                    "start_offset": start,
                    "end_offset": end,
                    "original_text": fragment,
                    "visual_id": element_id,
                }
            else:
                item = {
                    "kind": "table_cell_paragraph_span",
                    "table_index": int(location["table_index"]),
                    "row_index": int(location["row_index"]),
                    "col_index": int(location["col_index"]),
                    "cell_paragraph_index": int(location["cell_paragraph_index"]),
                    "start_offset": start,
                    "end_offset": end,
                    "original_text": fragment,
                    "visual_id": element_id,
                }
            sections.append({
                "section_type": "visual_text_span",
                "items": [item],
                "lines": [fragment],
                "word_count": word_count(fragment),
                "source_visual_id": element_id,
                "span_index": span_index,
            })
    return sections
