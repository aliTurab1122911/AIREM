from __future__ import annotations

import re
from collections import Counter
from copy import deepcopy
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CAPTION_RE = re.compile(r"^\s*(Figure|Fig\.?|Table)\s+(\d+)\s*([A-Za-z]?)\s*([:.\-–—])\s*(.+?)\s*$", re.I)


def _rgb(value: str | None, default: str = "000000") -> RGBColor:
    raw = re.sub(r"[^0-9A-Fa-f]", "", value or "")
    if len(raw) != 6:
        raw = default
    return RGBColor.from_string(raw.upper())


def _style_font_snapshot(style) -> Dict[str, Any]:
    font = style.font
    color = None
    try:
        color = str(font.color.rgb) if font.color and font.color.rgb else None
    except Exception:
        color = None
    return {
        "name": font.name,
        "size_pt": round(font.size.pt, 1) if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "color": color,
    }


def analyse_document_formatting(input_path: str | Path) -> Dict[str, Any]:
    doc = Document(str(input_path))
    fonts = Counter()
    sizes = Counter()
    colors = Counter()
    style_usage = Counter()
    manual_captions = []

    for paragraph in doc.paragraphs:
        style_usage[paragraph.style.name or "Normal"] += 1
        text = paragraph.text.strip()
        if CAPTION_RE.match(text):
            manual_captions.append(text[:200])
        for run in paragraph.runs:
            if run.font.name:
                fonts[run.font.name] += len(run.text) or 1
            if run.font.size:
                sizes[round(run.font.size.pt, 1)] += len(run.text) or 1
            try:
                if run.font.color and run.font.color.rgb:
                    colors[str(run.font.color.rgb)] += len(run.text) or 1
            except Exception:
                pass

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    style_usage[paragraph.style.name or "Normal"] += 1
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += len(run.text) or 1
                        if run.font.size:
                            sizes[round(run.font.size.pt, 1)] += len(run.text) or 1

    sections = []
    for index, section in enumerate(doc.sections, start=1):
        sections.append({
            "section": index,
            "width_inches": round(section.page_width.inches, 2),
            "height_inches": round(section.page_height.inches, 2),
            "orientation": "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait",
            "top_margin": round(section.top_margin.inches, 2),
            "bottom_margin": round(section.bottom_margin.inches, 2),
            "left_margin": round(section.left_margin.inches, 2),
            "right_margin": round(section.right_margin.inches, 2),
        })

    tracked_styles = {}
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        try:
            tracked_styles[style_name] = _style_font_snapshot(doc.styles[style_name])
        except KeyError:
            tracked_styles[style_name] = None

    return {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "top_fonts": fonts.most_common(8),
        "top_sizes": sizes.most_common(8),
        "top_colors": colors.most_common(8),
        "style_usage": style_usage.most_common(15),
        "styles": tracked_styles,
        "sections": sections,
        "manual_caption_count": len(manual_captions),
        "manual_caption_examples": manual_captions[:8],
    }


def _apply_run_typography(paragraph, font_name: str, size_pt: float, color: str, force_bold: bool | None = None) -> None:
    """Apply the visible typography selected in the studio to existing runs.

    Word documents often contain direct run formatting that overrides their
    paragraph style. Applying the same typography at run level keeps the final
    DOCX consistent with the browser live preview while preserving italic,
    underline and other emphasis unless a heading requires bold text.
    """
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
        run.font.size = Pt(float(size_pt))
        run.font.color.rgb = _rgb(color)
        if force_bold is True:
            run.bold = True


def _apply_document_run_typography(doc: Document, options: Dict[str, Any]) -> None:
    body_font = options.get("body_font") or "Arial"
    body_size = float(options.get("body_size") or 11)
    body_color = options.get("body_color") or "000000"
    heading_font = options.get("heading_font") or body_font
    heading_color = options.get("heading_color") or "1F4E79"
    caption_font = options.get("caption_font") or body_font
    caption_size = float(options.get("caption_size") or 10)
    caption_color = options.get("caption_color") or "404040"

    for paragraph in doc.paragraphs:
        style_name = (paragraph.style.name if paragraph.style else "Normal").strip().lower()
        if style_name == "title":
            _apply_run_typography(
                paragraph,
                options.get("title_font") or body_font,
                float(options.get("title_size") or 22),
                heading_color,
                True,
            )
        elif style_name.startswith("heading 1"):
            _apply_run_typography(paragraph, heading_font, float(options.get("h1_size") or 16), heading_color, True)
        elif style_name.startswith("heading 2"):
            _apply_run_typography(paragraph, heading_font, float(options.get("h2_size") or 14), heading_color, True)
        elif style_name.startswith("heading 3"):
            _apply_run_typography(paragraph, heading_font, float(options.get("h3_size") or 12), heading_color, True)
        elif style_name == "caption" or CAPTION_RE.match(paragraph.text.strip()):
            _apply_run_typography(paragraph, caption_font, caption_size, caption_color)
        else:
            _apply_run_typography(paragraph, body_font, body_size, body_color)


def _set_style(doc: Document, style_name: str, font_name: str | None, size_pt: float | None, color: str | None, bold: bool | None = None) -> None:
    try:
        style = doc.styles[style_name]
    except KeyError:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    if font_name:
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size_pt:
        style.font.size = Pt(float(size_pt))
    if color:
        style.font.color.rgb = _rgb(color)
    if bold is not None:
        style.font.bold = bool(bold)


def _add_field(paragraph, instruction: str, display: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)

    run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    paragraph.add_run(display)

    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _insert_before_first_paragraph(doc: Document, paragraphs) -> None:
    body = doc._element.body
    anchor = body[0] if len(body) else None
    for paragraph in paragraphs:
        if anchor is None:
            body.append(paragraph._p)
        else:
            anchor.addprevious(paragraph._p)


def _add_front_fields(doc: Document, add_toc: bool, add_lof: bool, add_lot: bool) -> None:
    created = []
    if add_toc:
        heading = doc.add_paragraph("Table of Contents")
        heading.style = "Heading 1" if "Heading 1" in [s.name for s in doc.styles] else "Normal"
        field = doc.add_paragraph()
        _add_field(field, 'TOC \\o "1-3" \\h \\z \\u', "Update field in Microsoft Word to generate the table of contents.")
        created.extend([heading, field])
    if add_lof:
        heading = doc.add_paragraph("List of Figures")
        heading.style = "Heading 1" if "Heading 1" in [s.name for s in doc.styles] else "Normal"
        field = doc.add_paragraph()
        _add_field(field, 'TOC \\h \\z \\c "Figure"', "Update field in Microsoft Word to generate the list of figures.")
        created.extend([heading, field])
    if add_lot:
        heading = doc.add_paragraph("List of Tables")
        heading.style = "Heading 1" if "Heading 1" in [s.name for s in doc.styles] else "Normal"
        field = doc.add_paragraph()
        _add_field(field, 'TOC \\h \\z \\c "Table"', "Update field in Microsoft Word to generate the list of tables.")
        created.extend([heading, field])
    if created:
        for paragraph in reversed(created):
            paragraph._element.getparent().remove(paragraph._element)
        _insert_before_first_paragraph(doc, created)


def _convert_captions(doc: Document) -> Dict[str, int]:
    counters = {"Figure": 0, "Table": 0}
    converted = {"Figure": 0, "Table": 0, "suffixed": 0}
    for paragraph in doc.paragraphs:
        match = CAPTION_RE.match(paragraph.text.strip())
        if not match:
            continue
        raw_label, raw_number, suffix, separator, title = match.groups()
        label = "Table" if raw_label.lower().startswith("table") else "Figure"
        manual_number = int(raw_number)
        if suffix:
            if manual_number == counters[label] + 1:
                counters[label] += 1
                instruction = f"SEQ {label} \\* ARABIC"
            elif manual_number == counters[label]:
                instruction = f"SEQ {label} \\c \\* ARABIC"
            else:
                counters[label] = manual_number
                instruction = f"SEQ {label} \\r {manual_number} \\* ARABIC"
            converted["suffixed"] += 1
        else:
            counters[label] += 1
            instruction = f"SEQ {label} \\* ARABIC"

        ppr = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
        for child in list(paragraph._p):
            paragraph._p.remove(child)
        if ppr is not None:
            paragraph._p.insert(0, ppr)
        try:
            paragraph.style = doc.styles["Caption"]
        except KeyError:
            pass
        paragraph.add_run(f"{label} ")
        _add_field(paragraph, instruction, str(manual_number))
        paragraph.add_run(f"{suffix}{separator} {title}")
        converted[label] += 1
    return converted


def _add_page_numbers(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = 2
        paragraph.add_run("Page ")
        _add_field(paragraph, "PAGE", "1")


def _set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")



def _set_table_borders(table, color: str = "808080", size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    clean_color = re.sub(r"[^0-9A-Fa-f]", "", color or "808080") or "808080"
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(max(2, min(int(size), 24))))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), clean_color.upper())
        borders.append(element)
    tbl_pr.append(borders)


def _set_page_border(section, color: str = "808080", size: int = 8) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgBorders"))
    if existing is not None:
        sect_pr.remove(existing)
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    clean_color = re.sub(r"[^0-9A-Fa-f]", "", color or "808080") or "808080"
    for edge in ["top", "left", "bottom", "right"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(max(2, min(int(size), 24))))
        element.set(qn("w:space"), "24")
        element.set(qn("w:color"), clean_color.upper())
        borders.append(element)
    sect_pr.append(borders)


def _set_header_footer_text(doc: Document, header_text: str, footer_text: str, different_first_page: bool = False) -> None:
    for section in doc.sections:
        section.different_first_page_header_footer = bool(different_first_page)
        if header_text:
            paragraph = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
            paragraph.text = header_text
            paragraph.alignment = 1
        if footer_text:
            paragraph = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
            paragraph.text = footer_text
            paragraph.alignment = 1


def apply_document_formatting(input_path: str | Path, output_path: str | Path, options: Dict[str, Any]) -> Dict[str, Any]:
    doc = Document(str(input_path))

    body_font = options.get("body_font") or "Arial"
    body_size = float(options.get("body_size") or 11)
    body_color = options.get("body_color") or "000000"
    _set_style(doc, "Normal", body_font, body_size, body_color)
    _set_style(doc, "Title", options.get("title_font") or body_font, float(options.get("title_size") or 22), options.get("heading_color") or "1F4E79", True)
    _set_style(doc, "Heading 1", options.get("heading_font") or body_font, float(options.get("h1_size") or 16), options.get("heading_color") or "1F4E79", True)
    _set_style(doc, "Heading 2", options.get("heading_font") or body_font, float(options.get("h2_size") or 14), options.get("heading_color") or "1F4E79", True)
    _set_style(doc, "Heading 3", options.get("heading_font") or body_font, float(options.get("h3_size") or 12), options.get("heading_color") or "1F4E79", True)
    _set_style(doc, "Caption", options.get("caption_font") or body_font, float(options.get("caption_size") or 10), options.get("caption_color") or "404040", False)
    _apply_document_run_typography(doc, options)

    line_spacing = float(options.get("line_spacing") or 1.15)
    space_after = float(options.get("space_after") or 6)
    for style_name in ["Normal", "Caption"]:
        style = doc.styles[style_name]
        style.paragraph_format.line_spacing = line_spacing
        style.paragraph_format.space_after = Pt(space_after)

    page_size = (options.get("page_size") or "a4").lower()
    orientation = (options.get("orientation") or "portrait").lower()
    for section in doc.sections:
        if page_size == "letter":
            width, height = Inches(8.5), Inches(11)
        else:
            width, height = Inches(8.27), Inches(11.69)
        if orientation == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = height, width
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = width, height
        section.top_margin = Inches(float(options.get("margin_top") or 1.0))
        section.bottom_margin = Inches(float(options.get("margin_bottom") or 1.0))
        section.left_margin = Inches(float(options.get("margin_left") or 1.0))
        section.right_margin = Inches(float(options.get("margin_right") or 1.0))
        if options.get("page_border"):
            _set_page_border(
                section,
                color=str(options.get("page_border_color") or "808080"),
                size=int(float(options.get("page_border_size") or 8)),
            )

    table_font = options.get("table_font") or body_font
    table_size = float(options.get("table_size") or 10)
    for table in doc.tables:
        if options.get("table_style"):
            try:
                table.style = options["table_style"]
            except Exception:
                pass
        _set_table_borders(
            table,
            color=str(options.get("table_border_color") or "808080"),
            size=int(float(options.get("table_border_size") or 4)),
        )
        for row_index, row in enumerate(table.rows):
            if row_index == 0 and options.get("repeat_header"):
                tr_pr = row._tr.get_or_add_trPr()
                tbl_header = OxmlElement("w:tblHeader")
                tbl_header.set(qn("w:val"), "true")
                tr_pr.append(tbl_header)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = table_font
                        run.font.size = Pt(table_size)
                        if row_index == 0 and options.get("bold_table_header"):
                            run.bold = True

    _set_header_footer_text(
        doc,
        str(options.get("header_text") or "").strip(),
        str(options.get("footer_text") or "").strip(),
        bool(options.get("different_first_page")),
    )

    caption_report = {"Figure": 0, "Table": 0, "suffixed": 0}
    if options.get("convert_captions"):
        caption_report = _convert_captions(doc)
    if options.get("page_numbers"):
        _add_page_numbers(doc)
    _add_front_fields(doc, bool(options.get("add_toc")), bool(options.get("add_lof")), bool(options.get("add_lot")))
    _set_update_fields(doc)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return {
        "output_path": str(output_path),
        "captions": caption_report,
        "page_numbers_added": bool(options.get("page_numbers")),
        "toc_added": bool(options.get("add_toc")),
        "list_of_figures_added": bool(options.get("add_lof")),
        "list_of_tables_added": bool(options.get("add_lot")),
        "sections_formatted": len(doc.sections),
        "tables_formatted": len(doc.tables),
        "page_border_added": bool(options.get("page_border")),
        "header_text_added": bool(str(options.get("header_text") or "").strip()),
        "footer_text_added": bool(str(options.get("footer_text") or "").strip()),
        "direct_run_typography_applied": True,
    }
