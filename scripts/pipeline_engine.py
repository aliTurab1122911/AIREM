from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Tuple

from docx import Document

from .docx_pipeline_common import (
    SEC,
    cell_text,
    is_caption_or_label,
    is_heading,
    iter_block_items,
    load_map,
    norm_text,
    para_text,
    parse_extract_text,
    replace_cell_text,
    replace_paragraph_text,
    replace_paragraph_span,
    save_json,
    section_word_count,
    sections_to_text,
    write_extract_docx,
)

SKIP_TABLE_IF_ANY_HEADER_CONTAINS = {"SUBNET", "GATEWAY"}


def should_skip_table(headers: List[str]) -> bool:
    joined = " | ".join(headers).upper()
    return any(token in joined for token in SKIP_TABLE_IF_ANY_HEADER_CONTAINS)


def _compact_heading_text(text: str) -> str:
    return " ".join((text or "").strip().lower().replace("_", " ").split())


START_EXTRACTION_HEADINGS = {
    "executive summary",
    "exec summary",
    "abstract",
    "introduction",
    "1 introduction",
    "1. introduction",
    "chapter 1 introduction",
}

END_EXTRACTION_HEADINGS = {
    "references",
    "reference",
    "reference list",
    "bibliography",
    "works cited",
    "appendix",
    "appendices",
}

FRONT_MATTER_HEADINGS = {
    "table of contents",
    "contents",
    "list of figures",
    "list of tables",
    "list of abbreviations",
    "abbreviations",
    "acknowledgements",
    "acknowledgments",
    "declaration",
}


def _is_start_heading(text: str) -> bool:
    t = _compact_heading_text(text)
    if not t:
        return False
    if t in START_EXTRACTION_HEADINGS:
        return True
    # Supports common numbering: "1.0 Introduction", "2 Executive Summary", etc.
    stripped = t.lstrip("0123456789.:- )(")
    return stripped in START_EXTRACTION_HEADINGS


def _is_end_heading(text: str) -> bool:
    t = _compact_heading_text(text)
    if not t:
        return False
    if t in END_EXTRACTION_HEADINGS:
        return True
    stripped = t.lstrip("0123456789.:- )(")
    return stripped in END_EXTRACTION_HEADINGS


def _is_front_matter_heading(text: str) -> bool:
    t = _compact_heading_text(text)
    if not t:
        return False
    if t in FRONT_MATTER_HEADINGS:
        return True
    stripped = t.lstrip("0123456789.:- )(")
    return stripped in FRONT_MATTER_HEADINGS


def extract_sections_from_docx(input_path: str | Path) -> List[Dict[str, Any]]:
    """Extract only main-body editable content from a DOCX.

    v8 clean-extraction rules:
    - skip cover page, table of contents, list of figures, list of tables, and other front matter
    - start extraction at Executive Summary, Abstract, or Introduction
    - stop before References, Bibliography, Works Cited, Appendix, or Appendices
    - extract normal paragraph text only; headings and captions are not extracted for paraphrasing
    - extract table body data only; table header rows are not extracted
    """
    src = Document(str(input_path))
    sections: List[Dict[str, Any]] = []
    current_paragraph_targets: List[Dict[str, Any]] = []
    current_paragraph_lines: List[str] = []
    body_para_index = -1
    table_index = -1
    in_main_body = False
    reached_end = False

    def flush_paragraph_group():
        nonlocal current_paragraph_targets, current_paragraph_lines
        if current_paragraph_lines:
            sections.append({
                "section_type": "paragraph_group",
                "items": current_paragraph_targets,
                "lines": current_paragraph_lines,
                "word_count": section_word_count(current_paragraph_lines),
            })
            current_paragraph_targets = []
            current_paragraph_lines = []

    for block in iter_block_items(src):
        if reached_end:
            break

        if block.__class__.__name__ == "Paragraph":
            body_para_index += 1
            text = para_text(block)
            clean = _compact_heading_text(text)

            # Never extract front matter headings or captions/lists.
            if _is_front_matter_heading(text):
                flush_paragraph_group()
                continue

            # Begin only when the document reaches the actual body.
            if not in_main_body:
                if _is_start_heading(text):
                    in_main_body = True
                # Do not extract the start heading itself. Extraction begins after it.
                continue

            # Stop before references/bibliography/appendices. Do not extract that heading or anything after it.
            if _is_end_heading(text):
                flush_paragraph_group()
                reached_end = True
                break

            if not text:
                flush_paragraph_group()
                continue

            # Go back to the original simpler scope: do not extract headings or captions/labels.
            if is_heading(block) or is_caption_or_label(text):
                flush_paragraph_group()
                continue

            # Avoid obvious generated lists from front matter if formatting is inconsistent.
            if clean.startswith("figure ") or clean.startswith("table "):
                flush_paragraph_group()
                continue

            current_paragraph_lines.append(text)
            current_paragraph_targets.append({
                "kind": "paragraph",
                "body_paragraph_index": body_para_index,
                "original_text": text,
                "style": block.style.name,
            })
        else:
            flush_paragraph_group()
            table_index += 1
            if not in_main_body:
                continue

            table = block
            if not table.rows:
                continue

            headers = [cell_text(c) for c in table.rows[0].cells]
            if should_skip_table(headers):
                continue

            row_count = len(table.rows)
            col_count = len(table.columns)

            # Clean scope: table body data only. Row 0 is assumed to be a header row and is skipped.
            for row_idx in range(1, row_count):
                lines: List[str] = []
                items: List[Dict[str, Any]] = []
                for col_idx in range(col_count):
                    cell = table.cell(row_idx, col_idx)
                    text = cell_text(cell) or ""
                    lines.append(text)
                    items.append({
                        "kind": "table_cell",
                        "table_index": table_index,
                        "row_index": row_idx,
                        "col_index": col_idx,
                        "original_text": text,
                    })
                if any(x.strip() for x in lines):
                    sections.append({
                        "section_type": "table_row",
                        "table_index": table_index,
                        "row_index": row_idx,
                        "headers": headers,
                        "items": items,
                        "lines": lines,
                        "word_count": section_word_count(lines),
                    })
    flush_paragraph_group()
    return sections

def split_sections_into_chunks(sections: List[Dict[str, Any]], max_words: int = 4500) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Split extraction sections into copy-safe chunks under max_words where possible.

    Splitting only happens between |sec| sections. If one section alone exceeds max_words, it is placed alone
    and reported as a warning because splitting inside a section would break reinsertion line matching.
    """
    chunks: List[Dict[str, Any]] = []
    warnings: List[Dict[str, Any]] = []
    cur_indices: List[int] = []
    cur_words = 0

    for idx, section in enumerate(sections):
        words = int(section.get("word_count") or section_word_count(section.get("lines", [])))
        if words > max_words:
            warnings.append({
                "section_number": idx + 1,
                "word_count": words,
                "message": "This single section exceeds the selected word limit and was placed in its own extract chunk.",
            })
        if cur_indices and cur_words + words > max_words:
            chunks.append({"section_indices": cur_indices, "word_count": cur_words})
            cur_indices = []
            cur_words = 0
        cur_indices.append(idx)
        cur_words += words
    if cur_indices:
        chunks.append({"section_indices": cur_indices, "word_count": cur_words})
    return chunks, warnings


def create_extraction_package(input_path: str | Path, output_dir: str | Path, map_path: str | Path, max_words: int = 4500) -> Dict[str, Any]:
    sections = extract_sections_from_docx(input_path)
    return create_extraction_package_from_sections(
        input_path=input_path,
        sections=sections,
        output_dir=output_dir,
        map_path=map_path,
        max_words=max_words,
        selection_metadata={"mode": "automatic_main_body"},
    )


def create_extraction_package_from_sections(
    input_path: str | Path,
    sections: List[Dict[str, Any]],
    output_dir: str | Path,
    map_path: str | Path,
    max_words: int = 4500,
    selection_metadata: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    """Create the normal extraction package from an explicit section selection."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    chunks, warnings = split_sections_into_chunks(sections, max_words=max_words)

    chunk_payloads: List[Dict[str, Any]] = []
    for chunk_number, chunk in enumerate(chunks, start=1):
        selected = [sections[i] for i in chunk["section_indices"]]
        selected_lines = [section["lines"] for section in selected]
        extract_name = f"2_extract_part_{chunk_number:02d}.docx"
        template_name = f"3_extract_airem_part_{chunk_number:02d}_paste_here.docx"
        extract_path = output_dir / extract_name
        template_path = output_dir / template_name
        write_extract_docx(selected_lines, extract_path)
        write_extract_docx(selected_lines, template_path)
        chunk_payloads.append({
            "chunk_number": chunk_number,
            "extract_file": extract_name,
            "edited_template_file": template_name,
            "section_indices": chunk["section_indices"],
            "section_count": len(chunk["section_indices"]),
            "word_count": chunk["word_count"],
            "text": sections_to_text(selected_lines),
        })

    slim_sections: List[Dict[str, Any]] = []
    for section in sections:
        slim_sections.append({key: value for key, value in section.items() if key != "lines"})

    payload = {
        "source_file": str(input_path),
        "separator": SEC,
        "max_words_per_extract": max_words,
        "section_count": len(sections),
        "chunk_count": len(chunks),
        "chunks": chunk_payloads,
        "sections": slim_sections,
        "warnings": warnings,
        "selection": selection_metadata or {"mode": "custom"},
        "notes": {
            "rule": "Every extract chunk starts and ends with |sec|. Keep all |sec| dividers in the edited/pasted text.",
            "matching": "Sections are matched by chunk order, section order, and line count inside each section.",
        },
    }
    save_json(payload, map_path)
    return payload


def validate_edited_chunks(mapping: Dict[str, Any], edited_texts: Dict[int, str]) -> Dict[str, Any]:
    issues: List[Dict[str, Any]] = []
    all_sections: List[List[str]] = []
    chunks_report: List[Dict[str, Any]] = []

    for chunk in mapping["chunks"]:
        n = int(chunk["chunk_number"])
        input_present = n in edited_texts
        raw_value = edited_texts.get(n, "")
        raw = "" if raw_value is None else str(raw_value)
        parsed = parse_extract_text(raw)
        normalised_lines = [norm_text(line) for line in raw.replace("\r", "").split("\n")]
        nonempty_lines = [line for line in normalised_lines if line]
        divider_count = sum(1 for line in nonempty_lines if line == SEC)
        expected_global_indices = chunk["section_indices"]
        expected_sections = [mapping["sections"][i] for i in expected_global_indices]
        chunk_issues: List[Dict[str, Any]] = []

        if not raw.strip():
            chunk_issues.append({
                "level": "error",
                "code": "empty_edited_chunk",
                "message": (
                    "Edited text is empty for this part. Run the rewrite first, or paste the complete edited "
                    "chunk before validation. This is not a list-formatting error."
                ),
                "expected_sections": len(expected_sections),
                "actual_sections": 0,
            })
        elif not parsed:
            chunk_issues.append({
                "level": "error",
                "code": "divider_only_chunk" if divider_count else "unparseable_edited_chunk",
                "message": (
                    "The edited text contains section dividers but no editable text. Paste or regenerate the full chunk."
                    if divider_count
                    else "The edited text could not be parsed into any sections. Paste or regenerate the full chunk."
                ),
                "expected_sections": len(expected_sections),
                "actual_sections": 0,
            })
        elif len(parsed) != len(expected_sections):
            no_dividers = divider_count == 0 and len(expected_sections) > 1
            chunk_issues.append({
                "level": "error",
                "code": "missing_section_dividers" if no_dividers else "section_count_mismatch",
                "message": (
                    "No |sec| divider lines were detected in this edited chunk. Use the app-generated rewrite "
                    "output or paste the complete marked chunk."
                    if no_dividers
                    else "Section count mismatch. Check missing or extra |sec| dividers in this pasted chunk."
                ),
                "expected_sections": len(expected_sections),
                "actual_sections": len(parsed),
            })
        for local_idx, expected in enumerate(expected_sections[:len(parsed)]):
            expected_items = expected["items"]
            actual_lines = parsed[local_idx]
            if len(actual_lines) != len(expected_items):
                chunk_issues.append({
                    "level": "error",
                    "code": "line_count_mismatch",
                    "section_number_in_chunk": local_idx + 1,
                    "global_section_number": expected_global_indices[local_idx] + 1,
                    "section_type": expected.get("section_type"),
                    "message": "Line count mismatch inside section. A paragraph/cell line was probably merged, deleted, or split.",
                    "expected_lines": len(expected_items),
                    "actual_lines": len(actual_lines),
                })
        chunks_report.append({
            "chunk_number": n,
            "valid": len(chunk_issues) == 0,
            "expected_sections": len(expected_sections),
            "actual_sections": len(parsed),
            "input_present": input_present,
            "input_character_count": len(raw),
            "nonempty_line_count": len(nonempty_lines),
            "divider_count": divider_count,
            "issue_count": len(chunk_issues),
            "issues": chunk_issues,
        })
        issues.extend({"chunk_number": n, **issue} for issue in chunk_issues)
        all_sections.extend(parsed)

    return {
        "valid": len(issues) == 0,
        "issue_count": len(issues),
        "section_count_expected": mapping["section_count"],
        "section_count_actual": len(all_sections),
        "chunks": chunks_report,
        "issues": issues,
        "parsed_sections": all_sections if len(issues) == 0 else [],
    }


def index_body_paragraphs_and_tables(doc: Document):
    paragraphs = []
    tables = []
    for block in iter_block_items(doc):
        if block.__class__.__name__ == "Paragraph":
            paragraphs.append(block)
        else:
            tables.append(block)
    return paragraphs, tables


def reinsert_sections(original_path: str | Path, mapping: Dict[str, Any], parsed_sections: List[List[str]], output_path: str | Path, log_path: str | Path | None = None) -> Dict[str, Any]:
    expected_sections = mapping["sections"]
    if len(parsed_sections) != len(expected_sections):
        raise ValueError(f"Section count mismatch: expected {len(expected_sections)}, got {len(parsed_sections)}")

    doc = Document(str(original_path))
    paragraphs, tables = index_body_paragraphs_and_tables(doc)
    replacements: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    span_edits: Dict[tuple, List[Dict[str, Any]]] = {}

    for sec_idx, (section, edited_lines) in enumerate(zip(expected_sections, parsed_sections), start=1):
        items = section["items"]
        if len(items) != len(edited_lines):
            errors.append({
                "section_number": sec_idx,
                "message": "Line count mismatch; section skipped.",
                "expected_lines": len(items),
                "actual_lines": len(edited_lines),
            })
            continue
        for item, new_text in zip(items, edited_lines):
            kind = item["kind"]
            if kind == "paragraph":
                idx = item["body_paragraph_index"]
                old = paragraphs[idx].text if idx < len(paragraphs) else ""
                replace_paragraph_text(paragraphs[idx], new_text)
                replacements.append({"section_number": sec_idx, "kind": kind, "body_paragraph_index": idx, "old_text": old, "new_text": new_text})
            elif kind == "table_cell":
                ti = item["table_index"]
                ri = item["row_index"]
                ci = item["col_index"]
                cell = tables[ti].cell(ri, ci)
                old = cell.text
                replace_cell_text(cell, new_text)
                replacements.append({"section_number": sec_idx, "kind": kind, "table_index": ti, "row_index": ri, "col_index": ci, "old_text": old, "new_text": new_text})
            elif kind == "paragraph_span":
                key = ("paragraph", int(item["body_paragraph_index"]))
                span_edits.setdefault(key, []).append({"section_number": sec_idx, "item": item, "new_text": new_text})
            elif kind == "table_cell_paragraph_span":
                key = (
                    "table_cell_paragraph", int(item["table_index"]), int(item["row_index"]),
                    int(item["col_index"]), int(item["cell_paragraph_index"]),
                )
                span_edits.setdefault(key, []).append({"section_number": sec_idx, "item": item, "new_text": new_text})

    for key, edits in span_edits.items():
        if key[0] == "paragraph":
            target = paragraphs[key[1]]
            location = {"body_paragraph_index": key[1]}
        else:
            _, ti, ri, ci, pi = key
            cell = tables[ti].cell(ri, ci)
            if pi >= len(cell.paragraphs):
                errors.append({"message": "Selected table-cell paragraph no longer exists.", "location": key})
                continue
            target = cell.paragraphs[pi]
            location = {"table_index": ti, "row_index": ri, "col_index": ci, "cell_paragraph_index": pi}

        original_target_text = target.text or ""
        for edit in sorted(edits, key=lambda value: int(value["item"]["start_offset"]), reverse=True):
            item = edit["item"]
            start_offset = int(item["start_offset"])
            end_offset = int(item["end_offset"])
            expected = str(item.get("original_text") or "")
            actual = (target.text or "")[start_offset:end_offset]
            if actual != expected:
                errors.append({
                    "section_number": edit["section_number"],
                    "message": "The selected text span no longer matches the uploaded document.",
                    "expected_text": expected,
                    "actual_text": actual,
                    **location,
                    "start_offset": start_offset,
                    "end_offset": end_offset,
                })
                continue
            replace_paragraph_span(target, start_offset, end_offset, edit["new_text"])
            replacements.append({
                "section_number": edit["section_number"],
                "kind": item["kind"],
                **location,
                "start_offset": start_offset,
                "end_offset": end_offset,
                "old_text": expected,
                "new_text": edit["new_text"],
            })

    if errors:
        result = {"valid": False, "errors": errors, "replacement_count": len(replacements)}
        if log_path:
            save_json(result, log_path)
        raise ValueError(f"Reinsertion stopped with {len(errors)} error(s).")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    result = {
        "valid": True,
        "original_file": str(original_path),
        "output_file": str(output_path),
        "replacement_count": len(replacements),
        "replacements": replacements,
    }
    if log_path:
        save_json(result, log_path)
    return result

